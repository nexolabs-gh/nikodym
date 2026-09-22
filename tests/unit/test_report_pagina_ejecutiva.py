"""Gates de la capa C1 de FLUJO-GUIADO-SCORECARD: la página ejecutiva del informe.

D-FLU-4 (el resumen final «es la página ejecutiva que ENTREGABLES-LEGIBLES pedía, con la misma
fuente que el informe») y D-FLU-11 fila C. El informe abre, tras la portada, con un capítulo
condicional «Resumen de la corrida» que reproduce el resumen final —los dos estados, las cifras
clave, qué revisar, las decisiones humanas con su motivo y dónde queda cada archivo— desde los
MISMOS constructores que ``Scorecard.summary()`` y que la pestaña Resultados
(:mod:`nikodym.guided.summaries`). Sin decisiones humanas el capítulo dice lo que dice la pantalla
(``SIN_DECISIONES``); sin corrida (un bundle armado a mano) no hay capítulo; y un resumen que no
se pueda armar no tumba el informe: el capítulo dice por qué no hay resumen.
"""

from __future__ import annotations

import html as html_module
import importlib.util
import re
from pathlib import Path
from typing import Any

import numpy as np
import pytest
from _ui_f1 import write_stacked_behavior_parquet
from test_guided_summaries import _ofensores

from nikodym.guided import STAGE_LABELS, Scorecard
from nikodym.guided.summaries import SIN_DECISIONES
from nikodym.report.builder import ReportBuilder
from nikodym.report.config import ReportConfig
from nikodym.report.document import (
    CHAPTER_SPECS,
    EXECUTIVE_SUMMARY_ID,
    RESULT_DOMAINS,
)
from nikodym.report.renderer import HtmlReportRenderer
from nikodym.report.results import ReportInputBundle

MOTIVO = "dato no disponible en originación"
_HAS_DOCX = importlib.util.find_spec("docx") is not None


@pytest.fixture(autouse=True)
def _usar_fake_binning_process(fake_binning_process: object) -> None:
    del fake_binning_process


def _corrida(raiz: Path, *, name: str, decidir: bool) -> Scorecard:
    """Una corrida completa de la puerta guiada con el doble de OptBinning (alcance de módulo)."""
    from conftest import FakeBinningProcess

    import nikodym.binning.transformer as transformer_module

    with pytest.MonkeyPatch.context() as parche:
        parche.setenv("PYTHONHASHSEED", "0")
        parche.setattr(transformer_module, "_import_binning_process", lambda: FakeBinningProcess)
        fuente = raiz / "cartera.parquet"
        write_stacked_behavior_parquet(fuente, repeats=50)
        sc = Scorecard(
            fuente,
            target="bad_flag",
            id="loan_id",
            cohort="cohort",
            oot_cohorts=["oot"],
            name=name,
            run_dir=raiz / "corridas",
            min_iv=0.0,
            formats=["md", "docx"],
        )
        sc._echo = lambda _texto: None
        if decidir:
            sc.exclude("score", reason=MOTIVO)
        sc.run()
        assert sc.study.run_context.status == "done", sc.study.run_context.error
        return sc


@pytest.fixture(scope="module")
def corrida(tmp_path_factory: pytest.TempPathFactory) -> Scorecard:
    """Con una decisión humana antes de correr: el motivo tiene que llegar a la página."""
    return _corrida(tmp_path_factory.mktemp("pagina"), name="con_decision", decidir=True)


@pytest.fixture(scope="module")
def corrida_sin_decisiones(tmp_path_factory: pytest.TempPathFactory) -> Scorecard:
    return _corrida(tmp_path_factory.mktemp("pagina-sin"), name="sin_decision", decidir=False)


def _html(sc: Scorecard) -> str:
    resultado = sc.study.artifacts.get("report", "result")
    return Path(resultado.html_path).read_text(encoding="utf-8")


def _seccion(html: str, section_id: str) -> str:
    inicio = html.index(f'data-section-id="{section_id}"')
    inicio = html.rindex("<section", 0, inicio)
    return html[inicio : html.index("</section>", inicio)]


def _pagina(sc: Scorecard) -> str:
    return _seccion(_html(sc), EXECUTIVE_SUMMARY_ID)


# ─────────────────────────── el capítulo y su lugar en el documento ───────────────────────────


def test_el_capitulo_es_el_primero_del_documento_y_condicional_al_scorecard() -> None:
    """Va tras la portada (antes del índice), sin número, y sólo en corridas de scorecard: una
    corrida IFRS 9 sin scorecard no recibe el molde del scorecard (insumo de H2)."""
    spec = CHAPTER_SPECS[0]
    assert spec.id == EXECUTIVE_SUMMARY_ID
    assert spec.kind == "summary"
    assert spec.numbered is False
    assert set(spec.requires_any_domain) == set(RESULT_DOMAINS)


def test_el_informe_abre_con_la_pagina_tras_la_portada(corrida: Scorecard) -> None:
    html = _html(corrida)
    portada = html.index('class="cover"')
    pagina = html.index(f'data-section-id="{EXECUTIVE_SUMMARY_ID}"')
    resumen_ejecutivo = html.index('id="exec-summary"')
    indice = html.index('data-kind="toc"')
    assert portada < pagina < resumen_ejecutivo < indice
    seccion = _pagina(corrida)
    assert 'data-kind="summary"' in seccion
    assert "Resumen de la corrida" in seccion
    assert "<h2>Resumen de la corrida</h2>" in seccion  # sin número: es materia preliminar


def test_el_indice_lista_la_pagina_antes_del_resumen_ejecutivo(corrida: Scorecard) -> None:
    html = _html(corrida)
    indice = _seccion(html, "toc")
    assert indice.index("#section-executive_summary") < indice.index("#exec-summary")
    # Y el sidebar de pantalla también la enumera en su lugar.
    assert html.index('href="#section-executive_summary"') < html.index('href="#exec-summary"')


# ─────────────────────────── la misma fuente que summary() y la pantalla ───────────────────────


def test_la_pagina_reproduce_el_resumen_final_de_la_puerta_guiada(corrida: Scorecard) -> None:
    final = corrida.summary()
    seccion = _pagina(corrida)
    escapar = html_module.escape
    assert escapar(final.validation) in seccion
    assert len(final.figures) == 5
    for rotulo, valor in final.figures:
        assert escapar(rotulo) in seccion, rotulo
        assert escapar(valor) in seccion, (rotulo, valor)
    for alerta in final.review:
        assert escapar(alerta) in seccion, alerta
    for rotulo in ("Ejecución", "Validación técnica", "Cifras clave", "Qué revisar"):
        assert rotulo in seccion, rotulo
    assert "Decisiones humanas registradas" in seccion
    assert "Dónde quedó cada archivo" in seccion


def test_la_ejecucion_nombra_las_etapas_que_corrieron_antes_del_informe(corrida: Scorecard) -> None:
    """Al renderizarse el informe la corrida sigue en curso: la página no afirma «completada»
    —eso lo dice ``summary()`` al terminar— sino qué corrió sin fallos antes del informe."""
    seccion = _pagina(corrida)
    assert "completada" not in seccion
    for etapa in corrida.steps:
        if etapa == "report":
            continue
        assert STAGE_LABELS[etapa] in seccion, etapa
    assert "este informe es la última etapa" in seccion


def test_las_decisiones_humanas_llegan_con_su_motivo(corrida: Scorecard) -> None:
    final = corrida.summary()
    assert final.decisions == (f"exclude score — «{MOTIVO}»",)
    seccion = _pagina(corrida)
    for linea in final.decisions:
        assert html_module.escape(linea) in seccion, linea
    assert html_module.escape(SIN_DECISIONES) not in seccion


def test_sin_decisiones_humanas_la_pagina_dice_lo_que_dice_la_pantalla(
    corrida_sin_decisiones: Scorecard,
) -> None:
    seccion = _pagina(corrida_sin_decisiones)
    assert html_module.escape(SIN_DECISIONES) in seccion
    assert corrida_sin_decisiones.summary().decisions == ()


def test_la_pagina_dice_donde_queda_cada_archivo_sin_inventar_rutas(corrida: Scorecard) -> None:
    """El HTML se nombra por su ruta real; los formatos pedidos, por la suya; el registro de
    auditoría y la ficha, por su nombre en la carpeta de evidencia (el informe se escribe antes de
    que la corrida se consolide y no conoce esa carpeta)."""
    resultado = corrida.study.artifacts.get("report", "result")
    seccion = _pagina(corrida)
    assert html_module.escape(str(resultado.html_path)) in seccion
    assert html_module.escape(str(resultado.md_path)) in seccion
    assert "audit_trail.jsonl" in seccion
    assert "carpeta de evidencia de la corrida" in seccion
    # Sin `purpose=` no hay ficha, y la página no la promete.
    assert "Ficha del modelo" not in seccion
    assert "Carpeta del proyecto" not in seccion  # el informe no la conoce; no la inventa


def test_la_pagina_no_filtra_identificadores_del_motor(corrida: Scorecard) -> None:
    """El gate de códigos internos del informe, extendido a la página ejecutiva (D-SIM-5)."""
    seccion = _pagina(corrida)
    texto = html_module.unescape(re.sub(r"<[^>]+>", " ", seccion))
    assert _ofensores(texto) == []


def test_el_study_conserva_el_preambulo_que_declaro(corrida: Scorecard) -> None:
    """De aquí lee el informe las decisiones humanas con motivo: la misma fuente que el trail."""
    preambulo = corrida.study.preamble
    decisiones = [p for _paso, p in preambulo if p.get("regla") == "decision_del_usuario"]
    assert len(decisiones) == 1
    assert decisiones[0]["motivo"] == MOTIVO
    assert decisiones[0]["variables"] == ["score"]


def test_el_preambulo_es_un_snapshot_que_nadie_muta(corrida: Scorecard) -> None:
    """Pasada 2 de Codex sobre C1: el payload anida listas que el llamador conserva. Ni mutar el
    original después de correr ni mutar lo que devuelve `preamble` cambia el snapshot: lo que
    se emitió al trail y lo que el informe lee son lo mismo."""
    original = next(d for d in corrida._decisions if d["regla"] == "decision_del_usuario")
    original["variables"].append("intruso")
    original["valor"]["selection.force_exclude"].append("intruso")
    try:
        decision = next(
            p for _paso, p in corrida.study.preamble if p.get("regla") == "decision_del_usuario"
        )
        assert decision["variables"] == ["score"]
        assert decision["valor"] == {"selection.force_exclude": ["score"]}
        decision["variables"].append("otro")
        de_nuevo = next(
            p for _paso, p in corrida.study.preamble if p.get("regla") == "decision_del_usuario"
        )
        assert de_nuevo["variables"] == ["score"]
    finally:
        original["variables"].remove("intruso")
        original["valor"]["selection.force_exclude"].remove("intruso")


def test_un_informe_regenerado_desde_un_study_recargado_dice_las_mismas_decisiones(
    corrida: Scorecard,
) -> None:
    """Pasada 4 de Codex: el preámbulo se persiste con la corrida (`run_metadata.json`) y vuelve
    con `Study.load`, así que un informe regenerado desde el Study recargado dice las mismas
    decisiones humanas que el trail, la ficha y el informe original."""
    from nikodym.core.study import Study

    recargado = Study.load(corrida.project_dir / "run" / "study", trust=True)
    assert recargado.preamble == corrida.study.preamble
    bundle = ReportBuilder(corrida.config.report).collect(recargado)
    assert bundle.summary is not None and bundle.summary["error"] is None
    assert bundle.summary["final"]["decisions"] == list(corrida.summary().decisions)
    original = corrida.study.artifacts.get("report", "input_bundle")
    assert bundle.summary["final"]["decisions"] == original.summary["final"]["decisions"]


def test_el_preambulo_persistido_es_el_prefijo_que_llego_al_trail() -> None:
    """Pasada 5 de Codex: si el sink falla en el evento N, el trail trae un prefijo y
    `run_context.preamble` trae exactamente el mismo prefijo, nunca el preámbulo entero; y una
    corrida nueva sobre el mismo Study empieza sin el preámbulo de la anterior."""
    from nikodym.audit.exceptions import AuditError
    from nikodym.core.audit import AuditEvent
    from nikodym.core.config import NikodymConfig
    from nikodym.core.study import Study

    class SinkQueSeLlena:
        def __init__(self) -> None:
            self.events: list[AuditEvent] = []

        def emit(self, event: AuditEvent) -> None:
            if event.kind == "decision" and any(e.kind == "decision" for e in self.events):
                raise AuditError("disco lleno")
            self.events.append(event)

    study = Study(NikodymConfig(), apply_global_seed=False)
    sink = SinkQueSeLlena()
    study.set_audit_sink(sink)
    declaraciones = [
        (None, {"regla": "primera", "umbral": None, "valor": 1, "accion": "declarar"}),
        (None, {"regla": "segunda", "umbral": None, "valor": 2, "accion": "declarar"}),
    ]
    with pytest.raises(AuditError):
        study.run(preamble=declaraciones)
    assert study.run_context.status == "failed"
    en_el_trail = [e.payload["regla"] for e in sink.events if e.kind == "decision"]
    assert en_el_trail == ["primera"]
    assert [p["regla"] for _paso, p in study.preamble] == en_el_trail
    # Una corrida nueva sobre el mismo Study arranca sin el preámbulo anterior.
    limpio = Study(NikodymConfig(), apply_global_seed=False)
    limpio.run_context.preamble = (("x", {"regla": "vieja", "accion": "a", "umbral": None}),)
    limpio.run(preamble=[])
    assert limpio.preamble == ()


def test_el_registro_de_auditoria_se_nombra_sin_su_ruta_absoluta(tmp_path: Path) -> None:
    """La suite completa acusó que dos corridas del mismo config por la interfaz daban HTML
    distintos: la interfaz reserva el trail en una ruta provisional con un token por corrida
    (`audit.trail_filename` absoluto, `.trail-<token>.jsonl`) y la página la imprimía. `audit`
    es INFRA —no entra al `config_hash`— y lo que varía con ella no entra al documento: con una
    ruta absoluta la página no imprime ni la ruta ni el nombre."""
    import nikodym

    fuente = tmp_path / "cartera.parquet"
    write_stacked_behavior_parquet(fuente, repeats=50)
    sc = Scorecard(
        fuente,
        target="bad_flag",
        id="loan_id",
        cohort="cohort",
        oot_cohorts=["oot"],
        name="trail_absoluto",
        run_dir=tmp_path / "corridas",
        min_iv=0.0,
    )
    from nikodym.audit.config import AuditConfig

    provisional = tmp_path / ".trail-0123456789abcdef.jsonl"
    config = sc.config.model_copy(
        update={"audit": AuditConfig(enabled=True, trail_filename=str(provisional))}
    )
    study = nikodym.run(config, run_dir=tmp_path / "corrida")
    assert study.run_context.status == "done", study.run_context.error
    html = Path(study.artifacts.get("report", "result").html_path).read_text(encoding="utf-8")
    seccion = _pagina_de(html)
    assert str(provisional) not in seccion and "0123456789abcdef" not in seccion
    assert "Registro de auditoría" in seccion
    assert "en la ruta que fija la sección audit del config" in html_module.unescape(seccion)


def _pagina_de(html: str) -> str:
    return _seccion(html, EXECUTIVE_SUMMARY_ID)


def test_regenerar_un_informe_no_mezcla_las_rutas_del_anterior(corrida: Scorecard) -> None:
    """Pasada 6 de Codex: un `Study` recargado conserva el artefacto `report.result` del informe
    ANTERIOR. Al regenerar con otro `output_dir`, «Dónde quedó cada archivo» tiene que decir las
    rutas de ESTE informe, una vez por rótulo, no dos destinos incompatibles."""
    from nikodym.core.study import Study

    recargado = Study.load(corrida.project_dir / "run" / "study", trust=True)
    otro = corrida.config.report.model_copy(update={"output_dir": "otro-destino"})
    bundle = ReportBuilder(otro).collect(recargado)
    archivos = bundle.summary["final"]["files"]
    rotulos = [rotulo for rotulo, _ruta in archivos]
    assert len(rotulos) == len(set(rotulos)), rotulos
    por_rotulo = dict(archivos)
    assert por_rotulo["Informe HTML"].startswith("otro-destino")
    previo = corrida.study.artifacts.get("report", "result").html_path
    assert previo not in dict(archivos).values()


def test_la_pagina_advierte_que_las_rutas_son_las_de_escritura(corrida: Scorecard) -> None:
    """Pasada 2 de Codex sobre C1: la interfaz copia el informe a `runs/<run_id>/` y una corrida
    posterior reescribe la ruta compartida; la página dice que sus rutas son las de escritura."""
    seccion = _pagina(corrida)
    assert "Las rutas son las que el informe escribió al generarse" in seccion


def test_con_pasos_despues_del_informe_la_pagina_no_afirma_lo_que_no_corrio(
    tmp_path: Path,
) -> None:
    """Pasada 1 de Codex sobre C1: `run.steps` conserva el orden del usuario y la validación es
    un insumo opcional del informe, así que `[…, "report", "validation"]` es válido y el informe
    se escribe ANTES de la validación. La página no puede afirmar entonces que cierra la corrida
    ni que la validación «no está en el config»: dice qué queda por correr y que no lo refleja."""
    import nikodym
    from nikodym.core.config import RunConfig

    fuente = tmp_path / "cartera.parquet"
    write_stacked_behavior_parquet(fuente, repeats=50)
    sc = Scorecard(
        fuente,
        target="bad_flag",
        id="loan_id",
        cohort="cohort",
        oot_cohorts=["oot"],
        name="reordenada",
        run_dir=tmp_path / "corridas",
        min_iv=0.0,
    )
    pasos = [paso for paso in sc.steps if paso != "validation"]
    pasos.append("validation")  # la validación formal corre DESPUÉS del informe
    config = sc.config.model_copy(update={"run": RunConfig(steps=pasos)})
    study = nikodym.run(config, run_dir=tmp_path / "corrida")
    assert study.run_context.status == "done", study.run_context.error
    assert study.artifacts.has("validation", "card")  # al final SÍ corrió: el informe no lo vio
    html = Path(study.artifacts.get("report", "result").html_path).read_text(encoding="utf-8")
    seccion = _seccion(html, EXECUTIVE_SUMMARY_ID)
    assert "última etapa" not in seccion
    assert "quedan por correr" in seccion
    assert STAGE_LABELS["validation"] in seccion
    assert "no puede reflejarlas" in seccion
    assert "no había corrido al emitir este informe" in seccion
    assert "no está en el config" not in seccion


# ─────────────────────────── los tres formatos ───────────────────────────


def test_el_qmd_lleva_la_pagina_antes_del_resumen_ejecutivo(corrida: Scorecard) -> None:
    resultado = corrida.study.artifacts.get("report", "result")
    qmd = Path(resultado.md_path).read_text(encoding="utf-8")
    assert qmd.index("## Resumen de la corrida") < qmd.index("## Resumen ejecutivo")
    # Todo valor dinámico va como texto literal de pandoc (puntuación ASCII con barra): se compara
    # sin las barras.
    literal = qmd.replace("\\", "")
    for linea in corrida.summary().decisions:
        assert linea in literal, linea
    for rotulo, valor in corrida.summary().figures:
        assert rotulo in literal and valor in literal, (rotulo, valor)


def test_el_qmd_neutraliza_texto_hostil_en_la_pagina(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Pasada 5 de Codex: los estados, las cifras, los archivos y el motivo de un resumen que no
    se armó también pueden traer marcado (una ruta o un mensaje con etiquetas). En la fuente
    editable van como texto literal de pandoc, nunca como HTML crudo ni enlaces."""
    import test_report_step as step_tests

    import nikodym.guided.summaries as fuente
    from nikodym.report.step import ReportStep

    hostil = '<img src=x onerror="alert(1)"> [x](http://mal) **negrita** # titulo'

    def revienta(*_args: Any, **_kwargs: Any) -> Any:
        raise RuntimeError(hostil)

    monkeypatch.setattr(fuente, "build_stage_summary", revienta)
    cfg = ReportConfig(output_dir=str(tmp_path), formats=["md"])
    study = step_tests._study_with_report_artifacts(config=cfg)
    result = ReportStep.from_config(cfg).execute(study, np.random.default_rng(1))
    qmd = Path(result.md_path).read_text(encoding="utf-8")
    inicio = qmd.index("## Resumen de la corrida")
    pagina = qmd[inicio : qmd.index("\n## ", inicio + 1)]
    assert hostil not in pagina
    # Toda la puntuación va con barra: ninguna etiqueta, enlace ni énfasis queda sin escapar.
    for activo in ("<img", "[x](http", "**negrita**", "# titulo"):
        assert not re.search(r"(?<!\\)" + re.escape(activo), pagina), activo
    assert "\\<img src\\=x" in pagina
    assert hostil in pagina.replace("\\", "")  # el texto sigue ahí, letra a letra


@pytest.mark.skipif(not _HAS_DOCX, reason="requiere el extra docx (python-docx)")
def test_el_word_lleva_la_pagina_antes_del_resumen_ejecutivo(corrida: Scorecard) -> None:
    import docx

    resultado = corrida.study.artifacts.get("report", "result")
    word = docx.Document(str(resultado.docx_path))
    titulos = [p.text for p in word.paragraphs if p.style.name.startswith("Heading 1")]
    assert titulos.index("Resumen de la corrida") < titulos.index("Resumen ejecutivo")
    texto = "\n".join(p.text for p in word.paragraphs)
    for linea in corrida.summary().decisions:
        assert linea in texto, linea
    celdas = {celda.text for tabla in word.tables for fila in tabla.rows for celda in fila.cells}
    for rotulo, valor in corrida.summary().figures:
        assert rotulo in celdas and valor in celdas, (rotulo, valor)


# ─────────────────────────── condicionalidad y degradación ───────────────────────────


def _lineage() -> Any:
    from datetime import UTC, datetime

    from nikodym.core.lineage import LineageBundle

    return LineageBundle(
        git_sha="abc123",
        git_dirty=False,
        data_hash="data123456789abcdef",
        config_hash="cfg123456789abcdef",
        root_seed=42,
        uv_lock_hash="uv123",
        library_versions={"nikodym": "1.19.0"},
        determinism_caveats=[],
        created_at=datetime(2026, 9, 21, 9, 30, tzinfo=UTC),
        schema_version="1.0.0",
    )


def test_un_bundle_armado_a_mano_no_tiene_capitulo(tmp_path: Path) -> None:
    """Sin corrida no hay resumen que reproducir: el bundle no lo trae y el capítulo no se emite,
    así que los goldens del renderer sobre bundles sintéticos no se mueven."""
    cfg = ReportConfig(sections={"missing_policy": "skip"})
    bundle = ReportInputBundle(
        lineage=_lineage(),
        cards={"performance": {"summary": "performance-card"}},
        tables={},
        figures={},
        sections=(),
    )
    assert bundle.summary is None
    secciones = ReportBuilder(cfg).build_sections(bundle)
    assert EXECUTIVE_SUMMARY_ID not in [s.id for s in secciones]
    html = HtmlReportRenderer(cfg).render(bundle.model_copy(update={"sections": secciones}))
    assert 'data-kind="summary"' not in html


def test_una_corrida_ifrs9_sin_scorecard_no_recibe_el_molde_del_scorecard() -> None:
    cfg = ReportConfig(sections={"missing_policy": "skip"})
    bundle = ReportInputBundle(
        lineage=_lineage(),
        cards={"provisioning_ifrs9": {"total_ecl_reported": 1.0, "total_ead": 10.0}},
        tables={},
        figures={},
        sections=(),
        summary={"final": {"execution": "x"}, "stages": [], "error": None},
    )
    secciones = ReportBuilder(cfg).build_sections(bundle)
    assert EXECUTIVE_SUMMARY_ID not in [s.id for s in secciones]


def test_un_resumen_que_no_se_puede_armar_no_tumba_el_informe(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Misma política que la pantalla: el informe sale entero y el capítulo dice por qué no hay
    resumen, en vez de esconder el hueco o de perder el informe por una frase."""
    import test_report_step as step_tests

    import nikodym.guided.summaries as fuente
    from nikodym.report.step import ReportStep

    def revienta(*_args: Any, **_kwargs: Any) -> Any:
        raise RuntimeError("la card llegó sin filas")

    monkeypatch.setattr(fuente, "build_stage_summary", revienta)
    cfg = ReportConfig(output_dir=str(tmp_path))
    study = step_tests._study_with_report_artifacts(config=cfg)
    result = ReportStep.from_config(cfg).execute(study, np.random.default_rng(1))
    html = Path(result.html_path).read_text(encoding="utf-8")
    seccion = _seccion(html, EXECUTIVE_SUMMARY_ID)
    assert "no se pudo armar" in seccion
    assert "la card llegó sin filas" in seccion
    assert "Cifras clave" not in seccion
