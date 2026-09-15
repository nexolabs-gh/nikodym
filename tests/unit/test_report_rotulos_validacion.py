"""Las cuatro tablas ``validation.*`` del informe pintan palabras, no identificadores del DTO.

Copy autorizado por Cami el 2026-09-15 (fuera de la enmienda VALIDACION-COTEJADA §7, con OK
propio): la tabla «Validación formal · calibración» del HTML, del Word y del Markdown imprimía
``hosmer_lemeshow``, ``pass``, ``not_evaluable``, ``performance_artifact``… mientras el panel y la
guía ya los traducían con los mapas de fuente única de ``nikodym.validation.results``. Aquí se exige
que el renderer aplique esos mismos mapas —ninguna palabra nueva— **sólo** en esas cuatro claves de
tabla y por clave de tabla (mismo patrón que el filtro de columnas de auditoría), sin mover los
encabezados: las trece columnas literales que los gates de las capas A y B exigen siguen ahí.

Control negativo preespecificado: aplicar el mapa a una tabla ajena con una columna ``test`` →
rojo (``test_una_tabla_ajena_con_una_columna_test_se_pinta_cruda``).
"""

from __future__ import annotations

import importlib.util
import io
import re
from datetime import UTC, datetime

import pandas as pd
import pytest

from nikodym.core.lineage import LineageBundle
from nikodym.report.builder import ReportBuilder
from nikodym.report.config import ReportConfig
from nikodym.report.markdown import MarkdownReportRenderer
from nikodym.report.renderer import HtmlReportRenderer, _table_view
from nikodym.report.results import ReportInputBundle
from nikodym.validation.config import CalibrationValidationConfig, ValidationConfig
from nikodym.validation.evaluator import ValidationEvaluator
from nikodym.validation.results import ValidationResult

_HAS_DOCX = importlib.util.find_spec("docx") is not None

#: Las trece columnas que la tabla de calibración del documento pinta, en su orden (literal).
_COLUMNAS_CALIBRACION: tuple[str, ...] = (
    "partition",
    "test",
    "grade",
    "n",
    "observed_defaults",
    "expected_pd",
    "observed_dr",
    "statistic",
    "degrees_of_freedom",
    "p_value",
    "alpha",
    "decision",
    "traffic_light",
)
#: Identificadores del DTO que hasta hoy llegaban crudos a la tabla y ya no deben aparecer en ella.
_SLUGS: tuple[str, ...] = (
    "hosmer_lemeshow",
    "brier",
    "performance_artifact",
    "stability_artifact",
    "recomputed",
    "not_evaluable",
    "score_psi",
    "pd_psi",
    "dev_vs_holdout",
    "dev_vs_oot",
    "stable",
)


def _frame() -> pd.DataFrame:
    """Tres particiones con PD distintas y target moderado: HL y Brier con veredicto."""
    filas = []
    for particion in ("desarrollo", "holdout", "oot"):
        for i in range(120):
            filas.append(
                {
                    "partition": particion,
                    "pd_calibrated": 0.02 + 0.005 * (i % 60),
                    "target": 1 if i % 9 == 0 else 0,
                    "grade": "A" if i % 2 else "B",
                }
            )
    return pd.DataFrame(filas)


def _resultado() -> ValidationResult:
    cfg = ValidationConfig(
        families=("discrimination", "calibration", "stability"),
        calibration=CalibrationValidationConfig(
            hl_n_groups=5, min_rows_per_group=10, binomial_by_grade=True
        ),
    )
    performance = pd.DataFrame(
        {
            "partition": ["desarrollo", "holdout", "oot"],
            "n_total": [120, 120, 120],
            "n_bad": [14, 14, 14],
            "auc": [0.71, 0.69, 0.68],
            "gini": [0.42, 0.38, 0.36],
            "ks": [0.31, 0.29, 0.27],
            "status": ["ok", "ok", "ok"],
        }
    )
    stability = pd.DataFrame(
        {
            "metric": ["score_psi", "pd_psi"],
            "comparison": ["dev_vs_holdout", "dev_vs_oot"],
            "feature": ["score", "pd_calibrated"],
            "value": [0.012, 0.3],
        }
    )
    return ValidationEvaluator.from_config(cfg).validate(
        calibrated_pd=_frame(),
        performance_metrics=performance,
        stability_metrics=stability,
        model_ref="scorecard-rotulos",
    )


def _lineage() -> LineageBundle:
    return LineageBundle(
        git_sha="abc123",
        git_dirty=False,
        data_hash="data123456789abcdef",
        config_hash="cfg123456789abcdef",
        root_seed=42,
        uv_lock_hash="uv123",
        library_versions={"nikodym": "1.16.0"},
        determinism_caveats=[],
        created_at=datetime(2026, 9, 15, 12, 0, tzinfo=UTC),
        schema_version="1.0.0",
    )


def _bundle(
    result: ValidationResult, extra_tables: dict[str, pd.DataFrame] | None = None
) -> ReportInputBundle:
    cfg = ReportConfig(sections={"missing_policy": "skip"})
    bundle = ReportInputBundle(
        lineage=_lineage(),
        cards={"validation": result.card.model_dump(mode="python")},
        results={"validation": result},
        tables={
            "validation.discrimination": result.discrimination,
            "validation.calibration": result.calibration,
            "validation.stability": result.stability,
            **(extra_tables or {}),
        },
        figures={},
        sections=(),
    )
    return bundle.model_copy(update={"sections": ReportBuilder(cfg).build_sections(bundle)})


def _cfg() -> ReportConfig:
    return ReportConfig(sections={"missing_policy": "skip"})


def _tabla_html(html: str, key: str) -> str:
    patron = rf'<table[^>]*data-table-key="{re.escape(key)}"[^>]*>(.*?)</table>'
    tabla = re.search(patron, html, re.S)
    assert tabla is not None, f"el documento no trae la tabla {key}"
    return tabla.group(1)


def _celdas(tabla: str) -> list[str]:
    return re.findall(r"<td>(.*?)</td>", tabla, re.S)


def test_la_tabla_de_calibracion_del_html_pinta_palabras_y_conserva_los_trece_encabezados() -> None:
    result = _resultado()
    html = HtmlReportRenderer(_cfg()).render(_bundle(result))
    tabla = _tabla_html(html, "validation.calibration")
    assert re.findall(r"<th>(.*?)</th>", tabla) == list(_COLUMNAS_CALIBRACION)
    celdas = _celdas(tabla)
    assert "Hosmer-Lemeshow" in celdas
    assert "Puntaje de Brier" in celdas
    assert "Jeffreys" in celdas
    assert "Desarrollo" in celdas and "Fuera de tiempo (OOT)" in celdas
    assert {"Pasa", "Falla"} & set(celdas)
    assert "Sin veredicto" in celdas  # el Brier es un puntaje, no una prueba de pasa/falla
    assert {"Verde", "Ámbar", "Rojo"} & set(celdas)
    for slug in _SLUGS:
        assert slug not in celdas, slug


def test_las_otras_tres_tablas_de_validacion_tambien_pintan_palabras() -> None:
    result = _resultado()
    html = HtmlReportRenderer(_cfg()).render(_bundle(result))
    discriminacion = _celdas(_tabla_html(html, "validation.discrimination"))
    assert "Reusada de la etapa de desempeño" in discriminacion
    assert "Evaluada" in discriminacion
    assert "performance_artifact" not in discriminacion and "ok" not in discriminacion
    estabilidad = _celdas(_tabla_html(html, "validation.stability"))
    assert "Reusado de la etapa de estabilidad" in estabilidad
    assert "PSI del score" in estabilidad or "PSI de la PD" in estabilidad
    assert "Desarrollo vs. Holdout" in estabilidad
    assert "Estable" in estabilidad and "Redesarrollar" in estabilidad
    assert "Pasa" in estabilidad and "Falla" in estabilidad
    assert "stability_artifact" not in estabilidad and "redevelop" not in estabilidad


def test_el_valor_del_dto_no_cambia_solo_su_pintura() -> None:
    """La tabla tidy, el JSON y la card siguen llevando el identificador: es copy del documento."""
    result = _resultado()
    assert set(result.calibration["test"]) >= {"hosmer_lemeshow", "brier"}
    vista = _table_view("validation.calibration", result.calibration, max_rows=50)
    assert vista["columns"] == list(_COLUMNAS_CALIBRACION)
    assert any("Hosmer-Lemeshow" in fila for fila in vista["rows"])
    assert set(result.calibration["test"]) >= {"hosmer_lemeshow", "brier"}  # el frame no se tocó


def test_una_tabla_ajena_con_una_columna_test_se_pinta_cruda() -> None:
    """El mapa se indexa por clave de tabla: una variable del usuario llamada ``test`` con el valor
    ``brier`` sigue diciendo ``brier`` en otra tabla. Control negativo declarado: aplicar el mapa
    por nombre de columna en todas las tablas pone este test rojo."""
    ajena = pd.DataFrame({"test": ["brier", "pass"], "decision": ["pass", "fail"], "n": [1, 2]})
    vista = _table_view("binning.tables", ajena, max_rows=50)
    assert vista["rows"] == [("brier", "pass", "1"), ("pass", "fail", "2")]
    result = _resultado()
    html = HtmlReportRenderer(_cfg()).render(_bundle(result, {"correlations.matrix": ajena}))
    if 'data-table-key="correlations.matrix"' in html:
        assert "brier" in _celdas(_tabla_html(html, "correlations.matrix"))


def test_el_markdown_pinta_las_mismas_palabras() -> None:
    result = _resultado()
    markdown = MarkdownReportRenderer.from_config(_cfg()).render(_bundle(result))
    assert "Hosmer-Lemeshow" in markdown and "Puntaje de Brier" in markdown
    assert "| partition | test |" in markdown  # el encabezado sigue literal
    assert "| hosmer_lemeshow |" not in markdown
    assert "| performance_artifact |" not in markdown


@pytest.mark.skipif(not _HAS_DOCX, reason="requiere el extra docx (python-docx)")
def test_el_word_pinta_las_mismas_palabras_y_conserva_los_encabezados() -> None:
    import docx

    from nikodym.report.docx import DocxReportRenderer

    result = _resultado()
    payload = DocxReportRenderer.from_config(_cfg()).render(_bundle(result))
    word = docx.Document(io.BytesIO(payload))
    tabla = next(
        t for t in word.tables if [c.text for c in t.rows[0].cells] == list(_COLUMNAS_CALIBRACION)
    )
    celdas = [c.text for fila in tabla.rows[1:] for c in fila.cells]
    assert "Hosmer-Lemeshow" in celdas and "Puntaje de Brier" in celdas
    assert "Pasa" in celdas or "Falla" in celdas
    assert "hosmer_lemeshow" not in celdas and "not_evaluable" not in celdas
